from io import BytesIO

import pytest
from docx import Document

from proposal_rag.rfp.docx_export import markdown_to_docx

MARKDOWN = """# Technical Proposal

## 1. Summary

Texte avec **gras** et normal.

- premier point
- second point

1. étape un

| ID | Status |
|----|--------|
| REQ-001 | COMPLIANT |
"""


def load(content: bytes):
    return Document(BytesIO(content))


def test_headings_lists_and_table_are_converted():
    document = load(markdown_to_docx(MARKDOWN))

    styles = [p.style.name for p in document.paragraphs]
    texts = [p.text for p in document.paragraphs]

    assert "Technical Proposal" in texts
    assert "Title" in styles
    assert "Heading 2" in styles
    assert styles.count("List Bullet") == 2
    assert "List Number" in styles

    table = document.tables[0]
    assert table.cell(0, 0).text == "ID"
    assert table.cell(1, 1).text == "COMPLIANT"


def test_bold_markers_become_bold_runs():
    document = load(markdown_to_docx("Texte avec **gras** ici"))
    runs = document.paragraphs[0].runs

    assert [r.text for r in runs] == ["Texte avec ", "gras", " ici"]
    assert runs[1].bold is True


def test_empty_markdown_rejected():
    with pytest.raises(ValueError):
        markdown_to_docx("   ")
