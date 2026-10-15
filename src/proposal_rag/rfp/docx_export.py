"""Conversion d'une proposition Markdown en document Word (.docx)."""

import re
from io import BytesIO

from docx import Document

_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _add_inline(paragraph, text: str) -> None:
    """Ajouter du texte en gérant le **gras** Markdown."""

    for part in _BOLD.split(text):
        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _flush_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return

    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            cell = table.cell(row_index, col_index)
            text = row[col_index] if col_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            _add_inline(paragraph, text)

            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def markdown_to_docx(markdown: str, title: str | None = None) -> bytes:
    """Produire un .docx à partir d'un sous-ensemble de Markdown :
    titres, listes, tableaux et texte en gras."""

    if not markdown or not markdown.strip():
        raise ValueError("La proposition est vide.")

    document = Document()
    table_rows: list[list[str]] = []
    first_heading_used = title is not None

    if title:
        document.add_heading(title, level=0)

    for line in markdown.splitlines():
        stripped = line.strip()

        # Tableau Markdown : accumuler les lignes jusqu'à la fin.
        if stripped.startswith("|"):
            if not _TABLE_SEPARATOR.match(stripped):
                table_rows.append(_split_row(stripped))
            continue

        if table_rows:
            _flush_table(document, table_rows)
            table_rows = []

        if not stripped:
            continue

        heading = _HEADING.match(stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()

            # Le premier "# Titre" devient le titre du document.
            if level == 1 and not first_heading_used:
                document.add_heading(text, level=0)
                first_heading_used = True
            else:
                document.add_heading(text, level=min(level, 4))
            continue

        bullet = _BULLET.match(line)
        if bullet:
            _add_inline(
                document.add_paragraph(style="List Bullet"),
                bullet.group(1),
            )
            continue

        numbered = _NUMBERED.match(line)
        if numbered:
            _add_inline(
                document.add_paragraph(style="List Number"),
                numbered.group(1),
            )
            continue

        _add_inline(document.add_paragraph(), stripped)

    _flush_table(document, table_rows)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
